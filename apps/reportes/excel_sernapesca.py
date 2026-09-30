"""Formato del XLSX de apoyo. No realiza consultas ni modifica registros."""
from datetime import date, datetime
from decimal import Decimal
from io import BytesIO
from math import ceil

from django.utils import timezone
from openpyxl import Workbook
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

NOTA_ALCANCE = (
    "Este archivo consolida información registrada en Belliazzi para facilitar su revisión. "
    "No reemplaza las declaraciones oficiales realizadas ante Sernapesca."
)
NOTA_REFERENCIAS = (
    "Postproceso y mermas son referencias de los seguimientos completos relacionados, "
    "no valores atribuidos ni prorrateados al lote. Pueden repetirse entre lotes y no deben sumarse entre filas. "
    "La cantidad asignada al lote proviene exclusivamente de sus aportes registrados."
)

# clave, encabezado, ancho. Los cambios de columnas se concentran aquí.
COLUMNAS = {
    "Abastecimiento": (
        ("recepcion", "ID recepción", 14), ("fecha", "Fecha/hora recepción (Chile)", 23),
        ("folio", "Folio de origen", 28), ("tipo_origen", "Tipo de origen", 22),
        ("agente", "Código agente", 20), ("proveedor", "Proveedor", 30), ("especie", "Especie", 26),
        ("peso_origen", "Peso de origen (kg)", 19), ("peso_recibido", "Peso recibido en planta (kg)", 22),
        ("diferencia", "Diferencia: recibido − origen (kg)", 24),
        ("usuario", "Usuario que registró", 30), ("observaciones", "Observaciones", 48),
    ),
    "Producción": (
        ("fecha", "Fecha de elaboración", 20), ("lote", "Código de lote de producción", 28),
        ("folio", "Folio de origen", 36), ("especie", "Especie", 26), ("ruta", "Ruta de procesamiento", 28),
        ("seguimientos", "Seguimientos asociados (IDs)", 28),
        ("postproceso", "Postproceso conocido de seguimientos (referencia, kg)", 30),
        ("asignado", "Cantidad asignada al lote (kg)", 24),
        ("merma", "Merma registrada de seguimientos (referencia, kg)", 28),
        ("descarte", "Descarte registrado de seguimientos (referencia, kg)", 28),
        ("perdida", "Pérdida registrada de seguimientos (referencia, kg)", 28),
        ("cajas", "Cantidad de cajas creadas", 22),
        ("empacado", "Peso neto real empacado conocido (kg)", 26),
        ("sin_peso", "Cajas con peso no documentado", 25),
        ("vencimiento", "Fecha de vencimiento", 20), ("usuario", "Usuario que creó lote", 30),
        ("observaciones", "Observaciones", 48), ("notas", "Información histórica incompleta", 48),
    ),
    "Destino": (
        ("despacho", "ID despacho", 14), ("fecha", "Fecha/hora despacho (Chile)", 23),
        ("tipo_destino", "Tipo de destino", 20), ("destinatario", "Destinatario", 32),
        ("rut", "RUT destinatario", 22), ("pais", "País destino", 22),
        ("tipo_documento", "Tipo documento", 26), ("numero_documento", "Número documento", 24),
        ("fecha_documento", "Fecha documento", 20), ("caja", "ID caja", 14),
        ("lote", "Código lote de producción", 28), ("especie", "Especie", 26),
        ("peso", "Peso neto real caja (kg)", 22), ("folio", "Folio de origen", 36),
        ("usuario", "Usuario que registró despacho", 30), ("observaciones", "Observaciones", 48),
    ),
}


def _celda(celda, valor):
    if isinstance(valor, datetime):
        # Excel no admite zona horaria: quitarla solo después de convertir a Chile.
        valor = timezone.localtime(valor, timezone.get_default_timezone()).replace(tzinfo=None)
        celda.number_format = "dd/mm/yyyy hh:mm"
    elif isinstance(valor, date):
        celda.number_format = "dd/mm/yyyy"
    elif isinstance(valor, Decimal):
        celda.number_format = "0.00"
    elif isinstance(valor, int):
        celda.number_format = "0"
    if isinstance(valor, str):
        # Texto literal: conserva ceros, folios y signos sin ejecutar fórmulas.
        valor = ILLEGAL_CHARACTERS_RE.sub(lambda m: f"\\u{ord(m.group()):04x}", valor)
        celda.value = valor
        celda.data_type = "s"
        celda.number_format = "@"
    else:
        celda.value = valor
    celda.font = Font(name="Arial", size=10, color="0F172A")
    celda.alignment = Alignment(vertical="top", wrap_text=True,
                                 horizontal="left" if isinstance(valor, str) else "right")


def _hoja(wb, nombre, columnas, filas):
    hoja = wb.create_sheet(nombre)
    hoja.sheet_view.showGridLines = False
    hoja.freeze_panes = "C2" if len(columnas) > 2 else "A2"
    hoja.print_title_rows = "1:1"
    hoja.row_dimensions[1].height = 64 if len(columnas) > 2 else 26
    for numero, (_, titulo, ancho) in enumerate(columnas, 1):
        celda = hoja.cell(1, numero, titulo)
        celda.font = Font(name="Arial", size=10, bold=True, color="164E63")
        celda.fill = PatternFill("solid", fgColor="ECFEFF")
        celda.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        hoja.column_dimensions[get_column_letter(numero)].width = ancho
    for indice, fila in enumerate(filas, 2):
        lineas = 1
        for numero, (clave, _, ancho) in enumerate(columnas, 1):
            valor = fila.get(clave)
            _celda(hoja.cell(indice, numero), valor)
            if isinstance(valor, str):
                lineas = max(lineas, sum(max(1, ceil(len(linea) / max(1, ancho - 3))) for linea in valor.split("\n")))
        hoja.row_dimensions[indice].height = min(409, max(24, lineas * 15 + 8))
    hoja.auto_filter.ref = f"A1:{get_column_letter(len(columnas))}{max(1, hoja.max_row)}"
    return hoja


def construir_excel_sernapesca(*, datos, inicio, fin, usuario, generado):
    wb = Workbook()
    wb.remove(wb.active)
    wb.properties.title = "Reporte de apoyo a Sernapesca"
    wb.properties.creator = usuario.username
    nombre = usuario.get_full_name()
    resumen = [
        ("Reporte", "Reporte de apoyo a Sernapesca"),
        ("Alcance", NOTA_ALCANCE),
        ("Generado", generado),
        ("Generado por", f"{usuario.username} ({nombre})" if nombre else usuario.username),
        ("Desde", inicio if inicio else "Desde el primer registro disponible"),
        ("Hasta (inclusive)", fin),
        ("Zona horaria", str(timezone.get_default_timezone())),
        ("Filas Abastecimiento", len(datos["Abastecimiento"])),
        ("Filas Producción", len(datos["Producción"])),
        ("Filas Destino", len(datos["Destino"])),
        ("Estado", "No hay registros para el período seleccionado." if not any(datos.values()) else "Contiene registros del período seleccionado."),
        ("Abastecimiento", "Una fila por detalle de recepción. Filtrado por fecha/hora de recepción. Diferencia = recibido en planta menos origen."),
        ("Producción", "Una fila por lote, filtrado por fecha de elaboración. Incluye todos sus aportes y cajas documentados al generar el reporte, aunque se registraran después del período."),
        ("Referencias del proceso", NOTA_REFERENCIAS),
        ("Postproceso incompleto", "Se suman solo seguimientos con un único pesaje postproceso. Sin pesaje o con varios, se excluye su peso y se indica en Información histórica incompleta. Sin pesos conocidos se deja vacío."),
        ("Mermas", "Sumas de mermas registradas en los seguimientos relacionados, sin filtro temporal adicional. Cero significa que no hay kg registrados de ese tipo. Sin seguimientos documentados queda vacío."),
        ("Destino", "Una fila por caja despachada, filtrada por fecha/hora del despacho. El folio proviene de las relaciones de origen, nunca se infiere del código de lote."),
        ("Pesos y campos vacíos", "Solo peso neto real documentado. Campos faltantes quedan vacíos. Los kg conocidos empacados excluyen cajas sin peso, que se cuentan por separado. No se sustituyen por pesos nominales."),
        ("Orígenes históricos", "Si hay más de un origen se muestran todos con sus identificadores y folios. No se selecciona arbitrariamente un único folio."),
        ("Texto", "Los identificadores se conservan como texto literal. Los caracteres de control incompatibles con Excel se representan mediante su código Unicode."),
    ]
    _hoja(wb, "Resumen", (("dato", "Dato", 30), ("valor", "Valor", 110)),
          [dict(dato=k, valor=v) for k, v in resumen])
    for nombre_hoja, columnas in COLUMNAS.items():
        _hoja(wb, nombre_hoja, columnas, datos[nombre_hoja])
    salida = BytesIO()
    wb.save(salida)
    return salida.getvalue()
