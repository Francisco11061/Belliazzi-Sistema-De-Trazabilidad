from datetime import date, datetime, timedelta
from decimal import Decimal
from io import BytesIO
from unittest.mock import patch
from zoneinfo import ZoneInfo

from django.contrib.auth.models import AnonymousUser
from django.core.exceptions import PermissionDenied, ValidationError
from django.db import connection
from django.test import TestCase, override_settings
from django.test.utils import CaptureQueriesContext
from django.urls import reverse
from openpyxl import load_workbook

from apps.inventario.models import Despacho, DetalleDespacho
from apps.trazabilidad.models import Caja, ConsumoLote, DetalleRecepcion, MermaProceso, Pesaje, RutaProceso
from .excel_sernapesca import COLUMNAS, NOTA_ALCANCE
from .forms import ReporteSernapescaForm
from .selectors import obtener_abastecimiento_reporte, obtener_produccion_reporte, obtener_destino_reporte
from .services import generar_reporte_apoyo_sernapesca
from .test_dashboard import AHORA, DatosDashboard


@override_settings(PASSWORD_HASHERS=["django.contrib.auth.hashers.MD5PasswordHasher"])
class ReporteSernapescaTests(DatosDashboard, TestCase):
    def exportar(self, **kwargs):
        contenido, nombre = generar_reporte_apoyo_sernapesca(usuario=self.jefe, fecha_fin=date(2026, 9, 30), **kwargs)
        return load_workbook(BytesIO(contenido)), nombre

    def aporte(self, lote, partida=None, kg="30.25"):
        partida = partida or self.partida()
        ConsumoLote.objects.create(lote_produccion=lote, partida=partida, cantidad_kg_utilizada=Decimal(kg))
        return partida

    def test_permisos_pantalla_descarga_y_sidebar(self):
        for usuario, permitido in ((self.jefe, True), (self.tecnico, True), (self.encargada, False), (self.operaria, False)):
            self.client.force_login(usuario)
            for nombre in ("sernapesca", "descargar_sernapesca"):
                with self.subTest(usuario=usuario.username, ruta=nombre):
                    self.assertEqual(self.client.get(self.url(nombre)).status_code, 200 if permitido else 403)
            inicio = self.client.get(reverse("usuarios:inicio")).content.decode()
            self.assertEqual(f'href="{self.url("sernapesca")}"' in inicio, permitido)

    def test_anonimo_y_post(self):
        for nombre in ("sernapesca", "descargar_sernapesca"):
            self.assertEqual(self.client.get(self.url(nombre)).status_code, 302)
        self.client.force_login(self.jefe)
        self.assertEqual(self.client.post(self.url("sernapesca")).status_code, 405)
        self.assertEqual(self.client.post(self.url("descargar_sernapesca")).status_code, 405)

    def test_servicio_revalida_permisos(self):
        for usuario in (AnonymousUser(), self.encargada, self.operaria):
            with self.assertRaises(PermissionDenied):
                generar_reporte_apoyo_sernapesca(usuario=usuario)
        self.jefe.is_active = False
        self.jefe.save()
        with self.assertRaises(PermissionDenied):
            self.exportar()

    def test_cambio_contrasena_pendiente(self):
        self.jefe.cambio_contrasena_pendiente = True
        self.jefe.save()
        self.client.force_login(self.jefe)
        for nombre in ("sernapesca", "descargar_sernapesca"):
            self.assertRedirects(self.client.get(self.url(nombre)), reverse("usuarios:cambiar_contrasena"))
        with self.assertRaises(PermissionDenied):
            self.exportar()

    def test_fechas_opcionales_y_orden(self):
        with patch("apps.reportes.forms.timezone.localdate", return_value=date(2026, 9, 30)):
            for datos in ({}, {"fecha_inicio": "2026-09-01"}, {"fecha_fin": "2026-09-30"},
                          {"fecha_inicio": "2026-09-30", "fecha_fin": "2026-09-30"},
                          {"fecha_inicio": "2026-09-01", "fecha_fin": "2026-09-30"}):
                form = ReporteSernapescaForm(datos)
                self.assertTrue(form.is_valid(), form.errors)
                self.assertEqual(form.cleaned_data["fecha_fin"], date(2026, 9, 30))

    def test_fechas_invalidas_no_descargan(self):
        self.client.force_login(self.jefe)
        for datos in ({"fecha_inicio": "2026-10-01", "fecha_fin": "2026-09-30"},
                      {"fecha_fin": "../archivo"}, {"fecha_inicio": "no-es-fecha"},
                      {"fecha_fin": "2026-02-30"}, {"fecha_fin": "9999-12-31"}):
            for nombre in ("sernapesca", "descargar_sernapesca"):
                response = self.client.get(self.url(nombre), datos)
                self.assertEqual(response.status_code, 400)
                self.assertNotIn("Content-Disposition", response)
        with self.assertRaises(ValidationError):
            generar_reporte_apoyo_sernapesca(usuario=self.jefe, fecha_inicio="../archivo")

    def test_pantalla_y_urls_existentes(self):
        self.client.force_login(self.jefe)
        response = self.client.get(self.url("sernapesca"))
        self.assertContains(response, NOTA_ALCANCE)
        self.assertContains(response, "Descargar Excel")
        self.assertContains(response, 'aria-current="page"', count=1)
        self.assertEqual(self.url(), "/dashboard/")
        self.assertEqual(self.url("umbrales"), "/dashboard/umbrales/")

    def test_excel_vacio_y_encabezados(self):
        wb, nombre = self.exportar()
        self.assertEqual(wb.sheetnames, ["Resumen", *COLUMNAS])
        self.assertEqual(nombre, "reporte_apoyo_sernapesca_desde_inicio_2026-09-30.xlsx")
        for hoja, columnas in COLUMNAS.items():
            self.assertEqual(wb[hoja].max_row, 1)
            self.assertEqual([c.value for c in wb[hoja][1]], [c[1] for c in columnas])
            self.assertEqual(wb[hoja].freeze_panes, "C2")
            self.assertTrue(wb[hoja].auto_filter.ref)
            self.assertTrue(wb[hoja]["A1"].font.bold)

    def test_descarga_headers_y_nombre(self):
        self.client.force_login(self.jefe)
        response = self.client.get(self.url("descargar_sernapesca"), {"fecha_inicio": "2026-09-01", "fecha_fin": "2026-09-30"})
        self.assertEqual(response["Content-Type"], "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
        self.assertIn("reporte_apoyo_sernapesca_2026-09-01_2026-09-30.xlsx", response["Content-Disposition"])
        self.assertIn("no-store", response["Cache-Control"])
        self.assertEqual(load_workbook(BytesIO(response.content)).sheetnames, ["Resumen", *COLUMNAS])

    def test_abastecimiento_varias_especies_y_pesos(self):
        partida = self.partida()
        detalle = partida.detalle_recepcion
        origen = detalle.origen_sernapesca
        origen.folio_origen = "00017/antiguo"
        origen.save()
        DetalleRecepcion.objects.create(recepcion=detalle.recepcion, origen_sernapesca=origen,
                                       especie=self.otra_especie, peso_origen_kg="20.50", peso_recepcion_kg="19.25")
        filas = obtener_abastecimiento_reporte(None, AHORA.date())
        self.assertEqual(len(filas), 2)
        self.assertEqual(filas[1]["folio"], "00017/antiguo")
        self.assertEqual(filas[1]["especie"], self.otra_especie.nombre)
        self.assertEqual(filas[1]["peso_origen"], Decimal("20.50"))
        self.assertEqual(filas[1]["peso_recibido"], Decimal("19.25"))
        self.assertEqual(filas[1]["diferencia"], Decimal("-1.25"))
        self.assertEqual(filas[1]["usuario"], self.jefe.username)

    def test_dia_final_completo_chile(self):
        p = self.partida()
        recepcion = p.detalle_recepcion.recepcion
        zona = ZoneInfo("America/Santiago")
        for fecha in (date(2026, 9, 30), date(2026, 9, 6), date(2026, 4, 4)):
            recepcion.fecha_hora_recepcion = datetime.combine(fecha, datetime.min.time(), zona) + timedelta(hours=23, minutes=59)
            recepcion.save()
            self.assertEqual(len(obtener_abastecimiento_reporte(fecha, fecha)), 1)
            recepcion.fecha_hora_recepcion = datetime.combine(fecha + timedelta(days=1), datetime.min.time(), zona)
            recepcion.save()
            self.assertEqual(obtener_abastecimiento_reporte(fecha, fecha), [])

    def test_aportes_parciales_pesos_reales_y_mermas_referencia(self):
        lote = self.lote()
        p = self.aporte(lote, kg="100")
        self.aporte(self.lote(), p, "25")
        Pesaje.objects.create(partida=p, tipo=Pesaje.POSTPROCESO, peso_kg=125, fecha_hora_evento=AHORA, registrado_por=self.jefe)
        for tipo in MermaProceso.Tipo.values:
            MermaProceso.objects.create(partida=p, tipo=tipo, cantidad_kg="2.50", motivo="Prueba", fecha_hora_evento=AHORA, registrado_por=self.jefe)
        for peso in ("10.25", "20.50", None):
            Caja.objects.create(lote_produccion=lote, peso_neto_kg=peso, peso_total_kg=999,
                                fecha_armado=AHORA, registrado_por=self.jefe)
        filas = obtener_produccion_reporte(None, AHORA.date())
        self.assertEqual([f["asignado"] for f in filas], [Decimal("100"), Decimal("25")])
        self.assertEqual([f["postproceso"] for f in filas], [Decimal("125"), Decimal("125")])
        self.assertEqual(filas[0]["empacado"], Decimal("30.75"))
        self.assertEqual(filas[0]["cajas"], 3)
        self.assertEqual(filas[0]["sin_peso"], 1)
        for campo in ("merma", "descarte", "perdida"):
            self.assertEqual(filas[0][campo], Decimal("2.50"))
        self.assertIn("referencia", COLUMNAS["Producción"][8][1])

    def test_multiples_aportes_y_origenes_historicos(self):
        lote = self.lote()
        lote.codigo_lote = "HISTORICO/ABC"
        lote.save()
        p = self.aporte(lote)
        otro = self.aporte(lote, kg="15.75")
        filas = obtener_produccion_reporte(None, AHORA.date())
        self.assertEqual(filas[0]["asignado"], Decimal("46"))
        self.assertEqual(filas[0]["lote"], "HISTORICO/ABC")
        self.assertIn("Múltiples orígenes históricos", filas[0]["folio"])
        self.assertIn(str(p.pk), filas[0]["seguimientos"])
        self.assertIn(str(otro.pk), filas[0]["seguimientos"])
        otro.detalle_recepcion.origen_sernapesca = p.detalle_recepcion.origen_sernapesca
        otro.detalle_recepcion.save()
        self.assertEqual(obtener_produccion_reporte(None, AHORA.date())[0]["folio"], "PRUEBA")

    def test_pesajes_ambiguos_y_datos_faltantes(self):
        lote = self.lote()
        fila = obtener_produccion_reporte(None, AHORA.date())[0]
        self.assertIsNone(fila["postproceso"])
        self.assertIsNone(fila["merma"])
        self.assertEqual(fila["ruta"], "")
        p = self.aporte(lote)
        for peso in (100, 95):
            Pesaje.objects.create(partida=p, tipo=Pesaje.POSTPROCESO, peso_kg=peso, fecha_hora_evento=AHORA, registrado_por=self.jefe)
        fila = obtener_produccion_reporte(None, AHORA.date())[0]
        self.assertIsNone(fila["postproceso"])
        self.assertIn("ambiguos excluidos", fila["notas"])
        self.exportar()

    def test_destino_varias_cajas_documento_y_trazabilidad(self):
        d = Despacho.objects.create(fecha_hora_despacho=AHORA, registrado_por=self.jefe,
                                   tipo_destino="EXPORTACION", destino="Cliente", pais_destino="Perú",
                                   tipo_documento="Factura", numero_documento="00027", fecha_documento=AHORA.date())
        for especie, peso in ((self.especie, "12.34"), (self.otra_especie, None)):
            caja = self.caja(peso, especie=especie)
            self.aporte(caja.lote_produccion)
            DetalleDespacho.objects.create(despacho=d, caja=caja)
        filas = obtener_destino_reporte(None, AHORA.date())
        self.assertEqual(len(filas), 2)
        self.assertEqual(filas[0]["peso"], Decimal("12.34"))
        self.assertIsNone(filas[1]["peso"])
        self.assertEqual(filas[1]["especie"], self.otra_especie.nombre)
        for f in filas:
            self.assertEqual(f["folio"], "PRUEBA")
            self.assertEqual(f["destinatario"], "Cliente")
            self.assertEqual(f["pais"], "Perú")
            self.assertEqual(f["numero_documento"], "00027")
        d.tipo_destino, d.rut_destinatario = "NACIONAL", "12345678-5"
        d.save()
        self.assertEqual(obtener_destino_reporte(None, AHORA.date())[0]["rut"], "12345678-5")

    def test_filtros_independientes_por_evento(self):
        p = self.partida()
        lote = self.lote()
        self.aporte(lote, p)
        d = self.despacho(fecha=AHORA + timedelta(days=2))
        lote.fecha_elaboracion = date(2026, 10, 1)
        lote.save()
        self.assertEqual(len(obtener_abastecimiento_reporte(date(2026, 9, 30), date(2026, 9, 30))), 1)
        self.assertEqual(len(obtener_produccion_reporte(date(2026, 10, 1), date(2026, 10, 1))), 1)
        self.assertEqual(obtener_destino_reporte(None, date(2026, 10, 1)), [])
        self.assertEqual(obtener_destino_reporte(date(2026, 10, 2), date(2026, 10, 2))[0]["despacho"], d.pk)

    def test_tipos_excel_y_texto_literal(self):
        p = self.partida()
        origen = p.detalle_recepcion.origen_sernapesca
        origen.folio_origen = "=1+1"
        origen.save()
        wb, _ = self.exportar()
        hoja = wb["Abastecimiento"]
        self.assertEqual(hoja["C2"].value, "=1+1")
        self.assertEqual(hoja["C2"].data_type, "s")
        self.assertEqual(hoja["H2"].data_type, "n")
        self.assertEqual(hoja["H2"].number_format, "0.00")
        self.assertEqual(hoja["B2"].value, datetime(2026, 9, 30, 12))
        self.assertEqual(hoja["B2"].data_type, "d")
        self.assertFalse(any(c.data_type == "f" for h in wb for fila in h for c in fila))

    def test_consultas_constantes_y_solo_lectura(self):
        self.aporte(self.lote())
        with CaptureQueriesContext(connection) as primero:
            self.exportar()
        for _ in range(3):
            self.aporte(self.lote())
        with CaptureQueriesContext(connection) as segundo:
            self.exportar()
        self.assertEqual(len(primero), len(segundo))
        self.assertTrue(all(q["sql"].lstrip().upper().startswith("SELECT") for q in segundo))

    def test_ruta_y_codigo_con_ceros_se_conservan(self):
        lote = self.lote()
        lote.codigo_lote = "00005010128092026"
        lote.ruta_proceso = RutaProceso.objects.create(nombre="Ruta reporte", especie=self.especie)
        lote.save()
        wb, _ = self.exportar()
        self.assertEqual(wb["Producción"]["B2"].value, lote.codigo_lote)
        self.assertEqual(wb["Producción"]["B2"].data_type, "s")
        self.assertEqual(wb["Producción"]["B2"].number_format, "@")
        self.assertEqual(wb["Producción"]["E2"].value, "Ruta reporte")

    def test_destino_historico_sin_origen_ni_peso(self):
        self.despacho(pesos=(None,))
        wb, _ = self.exportar()
        hoja = wb["Destino"]
        self.assertEqual(hoja.max_row, 2)
        for celda in ("D2", "E2", "F2", "M2", "N2"):
            self.assertIsNone(hoja[celda].value)
